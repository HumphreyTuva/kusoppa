"""Excel/CSV import (preview first, then apply) and Excel export for member data."""
import csv
import hashlib
import io
import json
import re
from datetime import date, datetime, time

from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel

from .models import Membership, Profile, Semester, SiteSettings, membership_status

FIELDS = [("full_name", "Name"), ("reg_no", "Registration number"), ("phone", "Phone"), ("year_of_study", "Year of study"),
          ("course", "Course"), ("leadership_interest", "Leadership interest"), ("mpesa_code", "M-Pesa code"),
          ("amount", "Amount paid"), ("paid_on", "Date paid"), ("semester", "Semester"), ("status", "Payment status"),
          ("member_id", "Member ID")]
LABELS = {"full_name": "name", "phone": "phone", "year_of_study": "year", "course": "course",
          "leadership_interest": "leadership interest", "reg_no": "registration number"}
SYN = {
    "full_name": ["name", "full name", "names", "student name", "member name", "fullname"],
    "reg_no": ["reg no", "reg number", "registration number", "registration no", "reg", "regno", "admission number", "adm no", "student number"],
    "phone": ["phone", "phone number", "mobile", "telephone", "tel", "contact", "phone no", "cell"],
    "year_of_study": ["year", "year of study", "yr"],
    "course": ["course", "programme", "program", "degree"],
    "leadership_interest": ["leadership interest", "leadership", "leader"],
    "mpesa_code": ["mpesa code", "m pesa code", "mpesa", "transaction code", "transaction id", "mpesa ref", "payment code", "code", "mpesa transaction code"],
    "amount": ["amount", "fee", "amount paid", "payment"],
    "paid_on": ["date", "date paid", "payment date", "timestamp", "paid on", "registration date"],
    "semester": ["semester", "sem", "term", "period"],
    "status": ["status", "payment status"],
    "member_id": ["member id", "id", "kusoppa id"],
}
ACCEPT = {"", "approved", "paid", "active", "yes", "ok", "confirmed", "complete", "completed"}
LEAD = {"yes": "yes", "y": "yes", "interested": "yes", "maybe": "maybe", "maybe later": "maybe", "no": "no",
        "not now": "no", "not right now": "no", "not interested": "no"}


def norm(h):
    return re.sub(r"[^a-z0-9]+", " ", str(h or "").lower()).strip()


def guess_field(header):
    if str(header or "").strip().startswith("["):  # semester columns in our own export
        return None
    n = norm(header)
    if not n:
        return None
    for f, names in SYN.items():
        if n in names:
            return f
    best, blen = None, 0
    for f, names in SYN.items():
        for s in names:
            if len(s) >= 4 and s in n and len(s) > blen:
                best, blen = f, len(s)
    return best


def auto_mapping(headers):
    m = {}
    for i, h in enumerate(headers):
        f = guess_field(h)
        if f and f not in m:
            m[f] = i
    return m


def detect_header(rows):
    best, score = 0, 1
    for i, r in enumerate(rows[:15]):
        s = sum(1 for c in r if c not in (None, "") and guess_field(c))
        if s > score:
            best, score = i, s
    return best if score > 1 else -1


def read_sheet(upload, sheet=None):
    name = (upload.name or "").lower()
    raw = upload.read()
    if len(raw) > 5 * 1024 * 1024:
        raise ValueError("The file is larger than 5 MB.")
    if name.endswith(".csv"):
        return ["CSV"], "CSV", list(csv.reader(io.StringIO(raw.decode("utf-8-sig", errors="replace"))))
    if not name.endswith(".xlsx"):
        raise ValueError("Please upload an .xlsx or .csv file. Save older .xls files as .xlsx first.")
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise ValueError("That file could not be read as an Excel workbook.")
    ws = wb[sheet] if sheet in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = []
    for r in ws.iter_rows(values_only=True):
        rows.append(list(r))
        if len(rows) > 5000:
            raise ValueError("The sheet has more than 5,000 rows. Split it into smaller files.")
    return wb.sheetnames, ws.title, rows


def text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return re.sub(r"\s+", " ", str(v)).strip()


def clean_reg(v):
    return re.sub(r"\s+", "", text(v)).upper()


def clean_code(v):
    return re.sub(r"\s+", "", text(v)).upper()


def clean_phone(v):
    t = re.sub(r"[\s\-().]", "", text(v))
    if not t:
        return "", ""
    if t.startswith("+254"):
        t = "0" + t[4:]
    elif t.startswith("254") and len(t) == 12:
        t = "0" + t[3:]
    elif len(t) == 9 and t[0] in "17":  # Excel dropped the leading zero
        t = "0" + t
    return t, ("" if re.fullmatch(r"0\d{9}|\+\d{9,15}", t) else f"phone '{t}' looks wrong")


def unique_phone_profile(phone):
    if not phone:
        return None
    qs = list(Profile.objects.filter(phone=phone, role="member")[:2])
    return qs[0] if len(qs) == 1 else None


def parse_amount(v):
    t = re.sub(r"[^\d.]", "", text(v))
    try:
        return int(float(t)) if t else None
    except ValueError:
        return None


def parse_date(v):
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and 20000 < v < 80000:
        try:
            return from_excel(v).date()
        except Exception:
            return None
    t = text(v)
    for fmt in ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d %b %Y", "%d %B %Y"):
        try:
            return datetime.strptime(t, fmt).date()
        except ValueError:
            pass
    return None


def find_semester(value, sems):
    n = norm(value)
    if not n:
        return None
    for s in sems:
        if norm(s.name) == n:
            return s
    for s in sems:
        sn = norm(s.name)
        if n in sn or sn in n:
            return s
    return None


def plan(rows, header_idx, mapping, opts):
    """Works out what an import would do, without saving anything."""
    sems, st = list(Semester.objects.all()), SiteSettings.get()
    sid = str(opts.get("semester") or "")
    default = Semester.objects.filter(pk=int(sid)).first() if sid.isdigit() else Semester.current()
    items, virtual, codes_seen, pairs_seen, identity_counts = [], {}, set(), set(), {}
    for n, r in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        if not any(text(c) for c in r):
            continue

        def get(f, r=r):
            i = mapping.get(f)
            return r[i] if i is not None and i < len(r) else None

        it = {"row": n, "issues": [], "changes": [], "pay": None, "profile": None, "member": "new", "note": ""}
        err = lambda m, it=it: it["issues"].append(("error", m))
        warn = lambda m, it=it: it["issues"].append(("warn", m))
        name, reg = text(get("full_name")), clean_reg(get("reg_no"))
        phone, pw = clean_phone(get("phone"))
        year, course = text(get("year_of_study")), text(get("course"))
        lead_raw = text(get("leadership_interest")).lower()
        lead = LEAD.get(lead_raw, "")
        if lead_raw and not lead:
            warn(f"leadership answer '{lead_raw}' not understood, ignored")
        if pw:
            warn(pw)

        identity = json.dumps([norm(name), norm(year), norm(course), norm(lead)], separators=(",", ":"))
        occurrence = identity_counts.get(identity, 0)
        identity_counts[identity] = occurrence + 1
        source = json.dumps([str(opts.get("source_name", "")).strip().casefold(),
                             str(opts.get("source_sheet", "")).strip().casefold(), identity, occurrence],
                            ensure_ascii=False, separators=(",", ":"))
        import_key = hashlib.sha256(source.encode("utf-8")).hexdigest()
        prof, mid = None, text(get("member_id"))
        if mid.isdigit():
            prof = Profile.objects.filter(pk=int(mid), role="member").first()
        if not prof and reg:
            prof = Profile.objects.filter(reg_no=reg, role="member").first()
        if not prof and phone:
            prof = unique_phone_profile(phone)
        if not prof:
            prof = Profile.objects.filter(import_key=import_key, role="member").first()
        if not prof and not name:
            warn("missing name; the row will be kept as an unnamed historical record")
        it["name"], it["reg_no"] = name or (prof.full_name if prof else ""), reg or (prof.reg_no if prof else "")
        it["profile"] = prof
        if prof:
            it["member"] = "existing"
            candidates = [("full_name", name), ("phone", phone), ("year_of_study", year), ("course", course), ("leadership_interest", lead)]
            if reg and prof.reg_no != reg:
                taken = Profile.objects.filter(reg_no=reg).exclude(pk=prof.pk).exists()
                if opts.get("overwrite") and not taken:
                    candidates.append(("reg_no", reg))
                else:
                    warn("registration number differs from the record, kept the existing one")
            for f, new in candidates:
                old = getattr(prof, f) or ""
                if not new or new == old:
                    continue
                if not old or opts.get("overwrite"):
                    it["changes"].append((f, old, new))
                else:
                    warn(f"keeps existing {LABELS[f]} ({old}); the sheet has '{new}'")
        else:
            it["fields"] = dict(full_name=name or f"Unnamed row {n}", reg_no=reg or None, phone=phone, import_key=import_key,
                                year_of_study=year, course=course, leadership_interest=lead)
            fingerprint = reg or (f"phone:{phone}" if phone else f"row:{n}")
            it["fingerprint"] = fingerprint
            if fingerprint in virtual:
                it["member"] = "repeat"

        code, amount, paid_on = clean_code(get("mpesa_code")), parse_amount(get("amount")), parse_date(get("paid_on"))
        if get("paid_on") not in (None, "") and not paid_on:
            warn("date not understood, ignored")
        status = text(get("status")).lower()
        if (code or amount or opts.get("all_paid")) and not any(lv == "error" for lv, _ in it["issues"]):
            if status not in ACCEPT:
                it["note"] = f"payment not imported (status is '{status}')"
            else:
                sem, sval = default, text(get("semester"))
                if "semester" in mapping and sval:
                    sem = find_semester(sval, sems)
                    if not sem:
                        err(f"Semester '{sval}' was not found. Add it in Settings first")
                elif not sem:
                    err("There is no semester to attach the payment to")
                if sem:
                    if code:
                        dup = Membership.objects.filter(mpesa_code=code).select_related("member", "semester").first()
                        if dup and not (prof and dup.member_id == prof.id and dup.semester_id == sem.id):
                            err(f"M-Pesa code {code} is already used by {dup.member.full_name} ({dup.semester.name})")
                        elif dup:
                            it["note"] = "payment already recorded"
                        elif code in codes_seen:
                            err(f"M-Pesa code {code} appears twice in this file")
                    if not any(lv == "error" for lv, _ in it["issues"]) and not it["note"]:
                        pair = ((prof.pk if prof else it.get("fingerprint") or it["reg_no"] or f"row:{it['row']}"), sem.id)
                        if prof and Membership.objects.filter(member=prof, semester=sem, status__in=["pending", "approved"]).exists():
                            it["note"] = f"already has a membership for {sem.name}"
                        elif pair in pairs_seen:
                            it["note"] = "repeated row for the same member and semester"
                        else:
                            it["pay"] = {"semester": sem, "code": code, "amount": amount, "paid_on": paid_on}
                            pairs_seen.add(pair)
                            if code:
                                codes_seen.add(code)
        it["level"] = "error" if any(lv == "error" for lv, _ in it["issues"]) else ("warn" if it["issues"] else "ok")
        if not prof and not (name or reg or phone) and not it["pay"]:
            err("Missing member details")
            it["level"] = "error"
        if it["level"] != "error" and not prof:
            virtual[it["fingerprint"]] = True
        if it["level"] == "error":
            it["pay"] = None
        items.append(it)

    # first payment is the joining payment, later semesters are renewals
    by_member = {}
    for it in items:
        if it["pay"]:
            key = it["profile"].pk if it["profile"] else it.get("fingerprint") or it["reg_no"] or f"row:{it['row']}"
            by_member.setdefault(key, []).append(it)
    for its in by_member.values():
        prof = its[0]["profile"]
        db_starts = list(Membership.objects.filter(member=prof, status="approved").values_list("semester__starts_on", flat=True)) if prof else []
        for it in its:
            s = it["pay"]["semester"].starts_on
            others = db_starts + [o["pay"]["semester"].starts_on for o in its if o is not it]
            it["pay"]["kind"] = "renewal" if any(x < s for x in others) else "new"
            if it["pay"]["amount"] is None:
                it["pay"]["amount"] = st.fee_renewal if it["pay"]["kind"] == "renewal" else st.fee_new
    return items


def summarise(items):
    ok = [i for i in items if i["level"] != "error"]
    return {
        "rows": len(items), "problems": len(items) - len(ok), "warnings": sum(1 for i in ok if i["level"] == "warn"),
        "new_members": len({i.get("fingerprint") or i["reg_no"] or f"row:{i['row']}" for i in ok if i["profile"] is None}),
        "existing_members": len({i["profile"].pk for i in ok if i["profile"]}),
        "updated_members": len({i["profile"].pk for i in ok if i["profile"] and i["changes"]}),
        "payments": sum(1 for i in ok if i["pay"]),
        "renewals": sum(1 for i in ok if i["pay"] and i["pay"]["kind"] == "renewal"),
        "skipped_payments": sum(1 for i in ok if i["note"]),
    }


def row_json(it):
    p = it["pay"]
    return {"row": it["row"], "name": it.get("name", ""), "reg_no": it.get("reg_no", ""), "member": it["member"], "level": it["level"],
            "issues": [m for _, m in it["issues"]], "note": it["note"],
            "changes": [f"{LABELS[f]}: {old or 'empty'} to {new}" for f, old, new in it["changes"]],
            "payment": p["kind"] if p else "none", "semester": p["semester"].name if p else "",
            "amount": p["amount"] if p else None, "code": p["code"] if p else ""}


def apply_plan(items, user, batch):
    created = updated = pays = 0
    made, now = {}, timezone.now()
    for it in items:
        if it["level"] == "error":
            continue
        key = it.get("fingerprint") or it["reg_no"] or f"row:{it['row']}"
        prof = it["profile"] or made.get(key)
        if prof is None:
            prof = made[key] = Profile.objects.create(import_batch=batch, **it["fields"])
            created += 1
        elif it["changes"]:
            for f, _, new in it["changes"]:
                setattr(prof, f, new)
            prof.save()
            updated += 1
        p = it["pay"]
        if p:
            m = Membership.objects.create(member=prof, semester=p["semester"], kind=p["kind"], amount=p["amount"], status="approved",
                                          mpesa_code=p["code"], reviewed_by=user, reviewed_at=now, import_batch=batch)
            if p["paid_on"]:
                Membership.objects.filter(pk=m.pk).update(submitted_at=timezone.make_aware(datetime.combine(p["paid_on"], time(12, 0))))
            pays += 1
    return {"members_created": created, "members_updated": updated, "payments_created": pays}


# ---------- export ----------
HEAD_FILL, HEAD_FONT = PatternFill("solid", fgColor="0B4D3B"), Font(bold=True, color="FFFFFF")


def put(ws, r, c, v, fmt=None):
    cell = ws.cell(row=r, column=c)
    cell.value = v
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        cell.data_type = "s"  # never let a name or note become an Excel formula
    if fmt:
        cell.number_format = fmt
    return cell


def add_sheet(wb, title, headers, rows, formats=None):
    ws = wb.create_sheet(title)
    for c, h in enumerate(headers, 1):
        cell = put(ws, 1, c, h)
        cell.fill, cell.font, cell.alignment = HEAD_FILL, HEAD_FONT, Alignment(vertical="center")
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            put(ws, r, c, v, (formats or {}).get(c))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for c in range(1, len(headers) + 1):
        w = max([len(str(headers[c - 1]))] + [len(str(r[c - 1])) for r in rows[:500] if r[c - 1] is not None])
        ws.column_dimensions[get_column_letter(c)].width = min(max(w + 2, 10), 42)
    return ws


def build_workbook(role):
    """Members for every official; Payments and Summary only for admin and treasurer."""
    wb = Workbook()
    wb.remove(wb.active)
    sems = list(Semester.objects.order_by("starts_on"))
    cur = Semester.current()
    people = list(Profile.objects.filter(role="member").select_related("user").prefetch_related("memberships__semester").order_by("full_name"))
    rows = []
    for p in people:
        ms = list(p.memberships.all())
        cells = []
        for s in sems:
            m = next(iter(sorted([x for x in ms if x.semester_id == s.id], key=lambda x: x.submitted_at, reverse=True)), None)
            cells.append("" if not m else (m.amount if m.status == "approved" else m.status.title()))
        rows.append([p.id, p.full_name, p.reg_no or "", p.email, p.phone, p.year_of_study, p.course,
                     p.get_leadership_interest_display() if p.leadership_interest else "", membership_status(p, cur).title(),
                     timezone.localtime(p.created_at).date()] + cells)
    add_sheet(wb, "Members", ["Member ID", "Name", "Registration number", "Email", "Phone", "Year of study", "Course", "Leadership interest",
                              "Membership status", "Joined"] + [f"[{s.name}]" for s in sems], rows, {5: "@", 10: "dd mmm yyyy"})
    if role in ("admin", "treasurer"):
        pay = Membership.objects.select_related("member", "semester", "reviewed_by").order_by("submitted_at")
        add_sheet(wb, "Payments", ["Member ID", "Name", "Registration number", "Semester", "Type", "Amount", "M-Pesa code", "Date paid", "Status", "Approved by"],
                  [[m.member_id, m.member.full_name, m.member.reg_no or "", m.semester.name, m.get_kind_display(), m.amount, m.mpesa_code,
                    timezone.localtime(m.submitted_at).date(), m.get_status_display(), (m.reviewed_by.email if m.reviewed_by else "")] for m in pay],
                  {8: "dd mmm yyyy"})
        summary = []
        for s in reversed(sems):
            ok = [m for m in Membership.objects.filter(semester=s, status="approved")]
            summary.append([s.name, len(ok), sum(1 for m in ok if m.kind == "new"), sum(1 for m in ok if m.kind == "renewal"), sum(m.amount for m in ok)])
        add_sheet(wb, "Summary", ["Semester", "Approved payments", "New members", "Renewals", "Total collected (Ksh)"], summary)
    return wb
