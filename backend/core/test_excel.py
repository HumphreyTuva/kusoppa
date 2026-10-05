import io
from datetime import datetime, timedelta

from django.contrib.auth.models import User
from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework.test import APITestCase

from .models import ImportBatch, Membership, Profile, Semester, SiteSettings


def xlsx(rows, title="Sheet1", name="members.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    for r in rows:
        ws.append(r)
    b = io.BytesIO()
    wb.save(b)
    return SimpleUploadedFile(name, b.getvalue())


HEAD = ["Timestamp", "Name", "Year", "Reg No.", "Phone number", "M-Pesa code", "Amount"]
SHEET = [["KUSOPPA registration list"], [], HEAD,
         [datetime(2026, 9, 15, 10, 30), "Ann Active", "Year 2", "p101/0001/2024", 712345678, "SGH7K2L9QX", 150],
         ["", "Bob Nocode", "Year 3", "p101/0002/2024", "0722000111", "", ""]]


class ExcelFlow(APITestCase):
    def setUp(self):
        cache.clear()
        t = timezone.localdate()
        self.old = Semester.objects.create(name="Jan to Apr 2026", starts_on=t - timedelta(days=200), ends_on=t - timedelta(days=100))
        self.cur = Semester.objects.create(name="Sep to Dec 2026", starts_on=t - timedelta(days=10), ends_on=t + timedelta(days=80))
        self.admin = User.objects.create_superuser("admin@x.com", "admin@x.com", "adminpass123")
        self.treasurer = self.staff("treas@x.com", "treasurer")
        self.secretary = self.staff("sec@x.com", "secretary")
        self.client.force_authenticate(self.treasurer)

    def staff(self, email, role):
        u = User.objects.create_user(email, email, "staffpass123")
        Profile.objects.create(user=u, full_name=role.title(), role=role)
        return u

    def go(self, which, rows, **data):
        return self.client.post(f"/api/manage/import/{which}/", {"file": xlsx(rows), **data}, format="multipart")

    def test_preview_finds_headers_and_maps_columns(self):
        r = self.go("preview", SHEET)
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data["header_row"], 2)
        self.assertEqual(set(r.data["mapping"]), {"paid_on", "full_name", "year_of_study", "reg_no", "phone", "mpesa_code", "amount"})
        s = r.data["summary"]
        self.assertEqual((s["new_members"], s["payments"], s["problems"]), (2, 1, 0))
        self.assertEqual(Profile.objects.filter(role="member").count(), 0)  # a preview saves nothing

    def test_apply_cleans_data_and_is_safe_to_repeat(self):
        r = self.go("apply", SHEET)
        self.assertEqual(r.status_code, 200, r.data)
        ann = Profile.objects.get(reg_no="P101/0001/2024")
        self.assertEqual(ann.phone, "0712345678")  # Excel had dropped the zero
        m = ann.memberships.get()
        self.assertEqual((m.status, m.amount, m.kind, m.mpesa_code, m.import_batch_id), ("approved", 150, "new", "SGH7K2L9QX", r.data["batch"]))
        self.assertEqual(m.submitted_at.date(), datetime(2026, 9, 15).date())
        self.assertFalse(Profile.objects.get(reg_no="P101/0002/2024").memberships.exists())
        again = self.go("preview", SHEET).data["summary"]
        self.assertEqual((again["new_members"], again["payments"], again["skipped_payments"]), (0, 0, 1))
        self.assertEqual(self.go("apply", SHEET).status_code, 400)

    def test_all_paid_option_and_renewal_kinds(self):
        self.assertEqual(self.go("preview", SHEET, all_paid="true").data["summary"]["payments"], 2)
        rows = [HEAD[1:4] + ["Semester", "M-Pesa code", "Amount"],
                ["Ann", "Year 2", "P101/0001/2024", "Sep to Dec 2026", "BBBBBBBBB2", ""],
                ["Ann", "Year 2", "P101/0001/2024", "Jan to Apr 2026", "AAAAAAAAA1", 150]]
        self.assertEqual(self.go("apply", rows).status_code, 200)
        got = {m.semester.name: (m.kind, m.amount) for m in Membership.objects.all()}
        self.assertEqual(got, {"Jan to Apr 2026": ("new", 150), "Sep to Dec 2026": ("renewal", 50)})

    def test_problems_are_reported_not_imported(self):
        rows = [HEAD, ["", "A", "Y1", "P1", "0700000001", "DUPLICATE1", 150], ["", "B", "Y1", "P2", "0700000002", "DUPLICATE1", 150],
                ["", "C", "Y1", "", "0700000003", "", ""]]
        s = self.go("preview", rows).data
        self.assertEqual(s["summary"]["problems"], 1)
        self.assertEqual([r["level"] for r in s["rows"]], ["ok", "error", "ok"])
        rows = [HEAD[:4] + ["Semester"], ["", "A", "Y1", "P1", "Winter 1999"]]
        self.assertEqual(self.go("preview", rows, mapping='{"full_name":1,"reg_no":3,"semester":4}', all_paid="true").data["summary"]["problems"], 1)
        self.assertIn("blocking", self.go("preview", [["Foo", "Bar"], ["x", "y"]]).data)

    def test_missing_identifiers_and_headerless_sheets_are_kept(self):
        rows = [["Sara Nafula", "Ksh 150.00"], ["Ronny Ochube", "Ksh 150.00"]]
        mapping = '{"full_name":0,"amount":1}'
        s = self.go("preview", rows, mapping=mapping, all_paid="true").data
        self.assertEqual(s["header_row"], -1)
        self.assertEqual((s["summary"]["new_members"], s["summary"]["payments"], s["summary"]["problems"]), (2, 2, 0))
        r = self.go("apply", rows, mapping=mapping, all_paid="true")
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(Profile.objects.filter(role="member", reg_no__isnull=True).count(), 2)
        self.assertEqual(Membership.objects.count(), 2)
        self.assertEqual(self.go("preview", rows, mapping=mapping, all_paid="true").data["summary"]["new_members"], 0)
        self.assertEqual(self.go("apply", rows, mapping=mapping, all_paid="true").status_code, 400)
        self.assertEqual(Profile.objects.filter(role="member", reg_no__isnull=True).count(), 2)

    def test_identical_rows_without_identifiers_remain_distinct_and_repeatable(self):
        rows = [["Sara Nafula", "Ksh 150.00"], ["Sara Nafula", "Ksh 150.00"]]
        mapping = '{"full_name":0,"amount":1}'
        result = self.go("apply", rows, mapping=mapping, all_paid="true")
        self.assertEqual(result.status_code, 200, result.data)
        self.assertEqual(Profile.objects.filter(role="member").count(), 2)
        self.assertEqual(self.go("preview", rows, mapping=mapping, all_paid="true").data["summary"]["new_members"], 0)
        self.assertEqual(Profile.objects.filter(role="member").count(), 2)

    def test_phone_matches_existing_member_when_registration_is_missing(self):
        p = Profile.objects.create(full_name="Imported Member", phone="0712345678", role="member")
        rows = [HEAD, ["", "Imported Member", "Year 4", "", "0712345678", "AAAAAAAAA1", 50]]
        s = self.go("preview", rows).data
        self.assertEqual((s["summary"]["existing_members"], s["summary"]["new_members"], s["summary"]["payments"]), (1, 0, 1))
        self.go("apply", rows)
        self.assertEqual(Profile.objects.filter(phone="0712345678").count(), 1)
        self.assertEqual(p.memberships.count(), 1)

    def test_details_update_only_when_asked(self):
        self.go("apply", SHEET)
        edited = [HEAD, ["", "Ann Active", "Year 2", "P101/0001/2024", "0799000000", "", ""]]
        s = self.go("preview", edited).data
        self.assertEqual(s["summary"]["updated_members"], 0)
        self.assertEqual(s["rows"][0]["level"], "warn")
        self.go("apply", edited, overwrite="true")
        self.assertEqual(Profile.objects.get(reg_no="P101/0001/2024").phone, "0799000000")

    def test_undo_removes_the_batch(self):
        r = self.go("apply", SHEET)
        b = r.data["batch"]
        self.assertEqual(self.client.post(f"/api/manage/import/batches/{b}/undo/").status_code, 200)
        self.assertEqual((Profile.objects.filter(role="member").count(), Membership.objects.count()), (0, 0))
        self.assertTrue(ImportBatch.objects.get(pk=b).undone)
        self.assertEqual(self.client.post(f"/api/manage/import/batches/{b}/undo/").status_code, 400)
        self.assertEqual(self.client.get("/api/manage/import/batches/").data[0]["undone"], True)

    def test_only_admin_and_treasurer_can_import(self):
        for u in (self.secretary, None):
            self.client.force_authenticate(u)
            self.assertIn(self.go("preview", SHEET).status_code, (401, 403))
            self.assertIn(self.go("apply", SHEET).status_code, (401, 403))
        self.client.force_authenticate(self.admin)
        self.assertEqual(self.go("preview", SHEET).status_code, 200)
        bad = self.client.post("/api/manage/import/preview/", {"file": SimpleUploadedFile("x.pdf", b"nope")}, format="multipart")
        self.assertEqual(bad.status_code, 400)

    def test_export_sheets_by_role_and_round_trip(self):
        self.go("apply", SHEET)
        Profile.objects.filter(reg_no="P101/0002/2024").update(full_name="=HYPERLINK(\"http://evil\")")
        raw = self.client.get("/api/manage/export/")
        wb = load_workbook(io.BytesIO(raw.content))
        self.assertEqual(wb.sheetnames, ["Members", "Payments", "Summary"])
        names = {c.value: c for c in wb["Members"]["B"]}
        self.assertEqual(names['=HYPERLINK("http://evil")'].data_type, "s")  # stored as text, never a formula
        self.assertEqual(wb["Summary"]["E2"].value, 150)
        self.client.force_authenticate(self.secretary)
        self.assertEqual(load_workbook(io.BytesIO(self.client.get("/api/manage/export/").content)).sheetnames, ["Members"])
        self.client.force_authenticate(self.treasurer)
        for sheet in ("Members", "Payments"):  # exported files can be uploaded back without creating anything new
            up = SimpleUploadedFile("back.xlsx", raw.content)
            s = self.client.post("/api/manage/import/preview/", {"file": up, "sheet": sheet}, format="multipart").data
            self.assertEqual((s["summary"]["new_members"], s["summary"]["payments"], s["summary"]["updated_members"]), (0, 0, 0), sheet)
